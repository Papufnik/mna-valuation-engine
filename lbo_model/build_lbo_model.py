"""
build_lbo_model.py -- generates LBO_Model.xlsx, a fully formula-driven
acquisition-financing model for the SAME illustrative deal
mna_valuation_engine.py values from the sell side (recast SDE $169,975.60,
base-case 3.00x multiple -> $509,926.80 enterprise value).

WHY THIS CAPITAL STRUCTURE, NOT A TEXTBOOK MEGA-CAP LBO: a deal this size
(sub-$200K SDE) does not clear for a syndicated senior term loan + high
yield bond structure -- no institutional lender underwrites acquisition
debt at that scale. The capital stack a real buyer uses at this deal size
is SBA 7(a) acquisition financing + a seller note + a sponsor equity
injection (the SBA's own minimum equity requirement for a change-of-
ownership 7(a) loan is 10%, which is why the sponsor-equity assumption
below is set at exactly that floor). This is the standard structure
independent sponsors and search funds actually use, and matching it to
the real deal size demonstrates the underwriting judgment call, not just
template-following -- see README > Design Decisions.

The underlying mechanics (sources & uses, debt amortization, free cash
flow, exit waterfall, IRR/MOIC) are identical to a large-cap LBO; only the
capital structure and pricing assumptions are scaled to the actual deal.

Every number that isn't a formula is a labeled, sourced, or explicitly
"illustrative -- confirm before relying on this" assumption on the
Assumptions tab -- consistent with mna_valuation_engine.py's own verified/
unverified labeling discipline elsewhere in this repo.
"""

import os

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_FILE = os.path.join(OUT_DIR, "LBO_Model.xlsx")

FONT_NAME = "Arial"
BLUE = Font(name=FONT_NAME, color="0000FF")
BLACK = Font(name=FONT_NAME, color="000000")
BOLD = Font(name=FONT_NAME, bold=True)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
HEADER_FILL = PatternFill(start_color="8B3A4A", end_color="8B3A4A", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
SECTION_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

CUR = '$#,##0;($#,##0);"-"'
PCT = "0.0%"
MULT = "0.00x"

YEARS_COLS = ["C", "D", "E", "F", "G"]  # Year 1..5
ALL_COLS = ["B"] + YEARS_COLS  # Year 0..5


def set_cell(ws, ref, value, font=BLACK, fmt=None, bold=False, note=None, fill=None):
    cell = ws[ref]
    cell.value = value
    cell.font = Font(name=FONT_NAME, color=font.color, bold=bold) if font else Font(name=FONT_NAME, bold=bold)
    if fmt:
        cell.number_format = fmt
    if note:
        cell.comment = Comment(note, "build_lbo_model.py")
    if fill:
        cell.fill = fill
    return cell


def section(ws, ref, text):
    cell = ws[ref]
    cell.value = text
    cell.font = BOLD
    cell.fill = SECTION_FILL


def title(ws, text):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT


def widen(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    title(ws, "LBO / Acquisition Financing Model -- Key Assumptions")
    widen(ws, {"A": 46, "B": 16})

    section(ws, "A3", "Deal Assumptions")
    set_cell(ws, "A4", "Entry Year 0 SDE (recast, from mna_valuation_engine.py)")
    set_cell(ws, "B4", 169975.60, font=BLUE, fmt=CUR,
             note="Source: mna_valuation_engine.py demo output, recastSDE field. "
                   "Same illustrative deal this repo's sell-side valuation values.")
    set_cell(ws, "A5", "Entry Multiple (x Entry SDE)")
    set_cell(ws, "B5", 3.00, font=BLUE, fmt=MULT,
             note="Matches mna_valuation_engine.py's own base-case multiple -- the "
                  "buyer underwrites the same price the sell-side model produced.")
    set_cell(ws, "A6", "Purchase Price (Enterprise Value)", bold=True)
    set_cell(ws, "B6", "=B4*B5", fmt=CUR, bold=True)
    set_cell(ws, "A7", "Estimated Transaction & Closing Costs")
    set_cell(ws, "B7", 15000, font=BLUE, fmt=CUR,
             note="Illustrative -- SBA guaranty fee, legal, and diligence/QoE costs "
                  "for a deal this size. Confirm actual quotes before relying on this.")
    set_cell(ws, "A8", "Total Uses", bold=True)
    set_cell(ws, "B8", "=B6+B7", fmt=CUR, bold=True)

    section(ws, "A10", "Capital Structure (% of Total Uses)")
    set_cell(ws, "A11", "SBA 7(a) Term Loan %")
    set_cell(ws, "B11", 0.75, font=BLUE, fmt=PCT)
    set_cell(ws, "A12", "Seller Note %")
    set_cell(ws, "B12", 0.15, font=BLUE, fmt=PCT)
    set_cell(ws, "A13", "Sponsor Equity %")
    set_cell(ws, "B13", 0.10, font=BLUE, fmt=PCT,
             note="Set at the SBA's minimum required equity injection for a "
                  "change-of-ownership 7(a) loan -- illustrative, confirm the "
                  "current SBA SOP before relying on this for a real deal.")
    set_cell(ws, "A14", "Check (should equal 100.0%)")
    set_cell(ws, "B14", "=B11+B12+B13", fmt=PCT)

    section(ws, "A16", "Debt Terms")
    set_cell(ws, "A17", "SBA 7(a) Rate (illustrative variable)")
    set_cell(ws, "B17", 0.115, font=BLUE, fmt=PCT,
             note="Illustrative Prime + lender spread, Q3 2026 ballpark. Confirm "
                  "current WSJ Prime + spread before using in a real model.")
    set_cell(ws, "A18", "SBA 7(a) Term (years)")
    set_cell(ws, "B18", 10, font=BLUE, fmt="0")
    set_cell(ws, "A19", "Seller Note Rate")
    set_cell(ws, "B19", 0.06, font=BLUE, fmt=PCT)
    set_cell(ws, "A20", "Seller Note Standby Period (yrs, interest-only)")
    set_cell(ws, "B20", 2, font=BLUE, fmt="0")
    set_cell(ws, "A21", "Seller Note Total Term (years)")
    set_cell(ws, "B21", 5, font=BLUE, fmt="0")

    section(ws, "A23", "Operating Assumptions")
    set_cell(ws, "A24", "Year 0 (LTM) Revenue")
    set_cell(ws, "B24", 1400000, font=BLUE, fmt=CUR,
             note="Illustrative -- specialty retail/hybrid revenue consistent with "
                  "a ~12% SDE margin at this deal size.")
    set_cell(ws, "A25", "Year 0 SDE Margin")
    set_cell(ws, "B25", "=B4/B24", fmt=PCT)
    set_cell(ws, "A26", "Annual Revenue Growth")
    set_cell(ws, "B26", 0.04, font=BLUE, fmt=PCT)
    set_cell(ws, "A27", "Margin Improvement, Years 3-5 (bps/yr)")
    set_cell(ws, "B27", 0.005, font=BLUE, fmt=PCT,
             note="Illustrative post-close operational tightening -- the standard "
                  "small-cap PE value-creation lever (systems, purchasing, labor "
                  "scheduling). Not assumed in Years 1-2 (integration period).")
    set_cell(ws, "A28", "Capex (% of Revenue)")
    set_cell(ws, "B28", 0.015, font=BLUE, fmt=PCT)
    set_cell(ws, "A29", "Incremental NWC (% of Revenue Growth $)")
    set_cell(ws, "B29", 0.10, font=BLUE, fmt=PCT)

    section(ws, "A31", "Exit Assumptions")
    set_cell(ws, "A32", "Base Case Exit Year")
    set_cell(ws, "B32", 5, font=BLUE, fmt="0")
    set_cell(ws, "A33", "Base Case Exit Multiple (x Exit SDE)")
    set_cell(ws, "B33", 3.00, font=BLUE, fmt=MULT)

    return ws


def build_sources_uses(wb):
    ws = wb.create_sheet("Sources & Uses")
    title(ws, "Sources & Uses")
    widen(ws, {"A": 40, "B": 16})

    section(ws, "A3", "Uses")
    set_cell(ws, "A4", "Purchase Price (Enterprise Value)")
    set_cell(ws, "B4", "=Assumptions!B6", fmt=CUR)
    set_cell(ws, "A5", "Transaction & Closing Costs")
    set_cell(ws, "B5", "=Assumptions!B7", fmt=CUR)
    set_cell(ws, "A6", "Total Uses", bold=True)
    set_cell(ws, "B6", "=SUM(B4:B5)", fmt=CUR, bold=True)

    section(ws, "A8", "Sources")
    set_cell(ws, "A9", "SBA 7(a) Term Loan")
    set_cell(ws, "B9", "=$B$6*Assumptions!B11", fmt=CUR)
    set_cell(ws, "A10", "Seller Note")
    set_cell(ws, "B10", "=$B$6*Assumptions!B12", fmt=CUR)
    set_cell(ws, "A11", "Sponsor Equity")
    set_cell(ws, "B11", "=$B$6*Assumptions!B13", fmt=CUR)
    set_cell(ws, "A12", "Total Sources", bold=True)
    set_cell(ws, "B12", "=SUM(B9:B11)", fmt=CUR, bold=True)
    set_cell(ws, "A13", "Check (Sources - Uses, should be 0)")
    set_cell(ws, "B13", "=B12-B6", fmt=CUR)

    section(ws, "A15", "Underwriting Metrics")
    set_cell(ws, "A16", "Total Debt / Entry SDE")
    set_cell(ws, "B16", "=(B9+B10)/Assumptions!B4", fmt=MULT)
    set_cell(ws, "A17", "Sponsor Equity / Total Uses")
    set_cell(ws, "B17", "=B11/B6", fmt=PCT)

    return ws


def build_operating_model(wb):
    ws = wb.create_sheet("Operating Model")
    title(ws, "Operating Model")
    widen(ws, {"A": 40})
    for c in ALL_COLS:
        ws.column_dimensions[c].width = 14

    headers = ["", "Year 0 (LTM)", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
    for i, col in enumerate(["A"] + ALL_COLS):
        set_cell(ws, f"{col}3", headers[i], bold=True, fill=HEADER_FILL if col != "A" else None,
                  font=HEADER_FONT if col != "A" else BOLD)

    # hidden-ish helper row of numeric year indices, used by Returns Analysis
    # for dynamic INDEX/MATCH lookups keyed off the exit-year assumption.
    set_cell(ws, "A2", "Year index (helper row for lookups)")
    for i, col in enumerate(ALL_COLS):
        set_cell(ws, f"{col}2", i, font=BLUE, fmt="0")

    set_cell(ws, "A4", "Revenue")
    set_cell(ws, "B4", "=Assumptions!B24", fmt=CUR)
    for prev, col in zip(ALL_COLS, YEARS_COLS):
        set_cell(ws, f"{col}4", f"={prev}4*(1+Assumptions!$B$26)", fmt=CUR)

    set_cell(ws, "A5", "Revenue Growth %")
    for prev, col in zip(ALL_COLS, YEARS_COLS):
        set_cell(ws, f"{col}5", f"={col}4/{prev}4-1", fmt=PCT)

    set_cell(ws, "A6", "SDE / EBITDA Margin %")
    set_cell(ws, "B6", "=Assumptions!B25", fmt=PCT)
    set_cell(ws, "C6", "=B6", fmt=PCT)
    set_cell(ws, "D6", "=C6", fmt=PCT)
    set_cell(ws, "E6", "=D6+Assumptions!$B$27", fmt=PCT)
    set_cell(ws, "F6", "=E6+Assumptions!$B$27", fmt=PCT)
    set_cell(ws, "G6", "=F6+Assumptions!$B$27", fmt=PCT)

    set_cell(ws, "A7", "SDE / EBITDA ($)", bold=True)
    for col in ALL_COLS:
        set_cell(ws, f"{col}7", f"={col}4*{col}6", fmt=CUR, bold=True)

    set_cell(ws, "A9", "Less: Capex")
    for col in YEARS_COLS:
        set_cell(ws, f"{col}9", f"=-{col}4*Assumptions!$B$28", fmt=CUR)

    set_cell(ws, "A10", "Less: Incremental NWC")
    for prev, col in zip(ALL_COLS, YEARS_COLS):
        set_cell(ws, f"{col}10", f"=-({col}4-{prev}4)*Assumptions!$B$29", fmt=CUR)

    set_cell(ws, "A12", "Unlevered Free Cash Flow (pre-debt-service)", bold=True)
    for col in YEARS_COLS:
        set_cell(ws, f"{col}12", f"={col}7+{col}9+{col}10", fmt=CUR, bold=True)

    ws["A14"] = ("Note: modeled at the pre-tax, pass-through SDE level appropriate for "
                 "a small-business acquisition, not a full corporate-tax LBO -- see "
                 "README > Limitations.")
    ws["A14"].font = Font(name=FONT_NAME, italic=True, size=9, color="666666")

    return ws


def build_debt_schedule(wb):
    ws = wb.create_sheet("Debt Schedule")
    title(ws, "Debt Schedule")
    widen(ws, {"A": 44})
    for c in ALL_COLS:
        ws.column_dimensions[c].width = 14

    headers = ["", "At Close", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
    for i, col in enumerate(["A"] + ALL_COLS):
        set_cell(ws, f"{col}3", headers[i], bold=True, fill=HEADER_FILL if col != "A" else None,
                  font=HEADER_FONT if col != "A" else BOLD)

    # --- SBA 7(a) Term Loan ---
    section(ws, "A5", "SBA 7(a) Term Loan")
    set_cell(ws, "A6", "Beginning Balance")
    set_cell(ws, "B6", "='Sources & Uses'!B9", fmt=CUR)
    set_cell(ws, "C6", "=B6", fmt=CUR)
    set_cell(ws, "D6", "=C10", fmt=CUR)
    set_cell(ws, "E6", "=D10", fmt=CUR)
    set_cell(ws, "F6", "=E10", fmt=CUR)
    set_cell(ws, "G6", "=F10", fmt=CUR)

    set_cell(ws, "A7", "Annual Debt Service (level P&I)")
    for col in YEARS_COLS:
        set_cell(ws, f"{col}7",
                  "=-PMT(Assumptions!$B$17,Assumptions!$B$18,'Sources & Uses'!$B$9)",
                  fmt=CUR)

    set_cell(ws, "A8", "Interest Expense")
    for col in YEARS_COLS:
        set_cell(ws, f"{col}8", f"={col}6*Assumptions!$B$17", fmt=CUR)

    set_cell(ws, "A9", "Principal Payment")
    for col in YEARS_COLS:
        set_cell(ws, f"{col}9", f"={col}7-{col}8", fmt=CUR)

    set_cell(ws, "A10", "Ending Balance", bold=True)
    for col in YEARS_COLS:
        set_cell(ws, f"{col}10", f"={col}6-{col}9", fmt=CUR, bold=True)

    # --- Seller Note ---
    section(ws, "A13", "Seller Note (2-yr standby, then amortizes)")
    set_cell(ws, "A14", "Beginning Balance")
    set_cell(ws, "B14", "='Sources & Uses'!B10", fmt=CUR)
    set_cell(ws, "C14", "=B14", fmt=CUR)
    set_cell(ws, "D14", "=C18", fmt=CUR)
    set_cell(ws, "E14", "=D18", fmt=CUR)
    set_cell(ws, "F14", "=E18", fmt=CUR)
    set_cell(ws, "G14", "=F18", fmt=CUR)

    set_cell(ws, "A15", "Interest Expense")
    for col in YEARS_COLS:
        set_cell(ws, f"{col}15", f"={col}14*Assumptions!$B$19", fmt=CUR)

    set_cell(ws, "A16", "Post-Standby Annual P&I (Yrs 3-5)")
    set_cell(ws, "E16", "=-PMT(Assumptions!$B$19,Assumptions!$B$21-Assumptions!$B$20,E14)", fmt=CUR)
    set_cell(ws, "F16", "=E16", fmt=CUR)
    set_cell(ws, "G16", "=F16", fmt=CUR)

    set_cell(ws, "A17", "Principal Payment")
    set_cell(ws, "C17", 0, fmt=CUR, note="Standby period -- interest-only, Years 1-2.")
    set_cell(ws, "D17", 0, fmt=CUR)
    set_cell(ws, "E17", "=E16-E15", fmt=CUR)
    set_cell(ws, "F17", "=F16-F15", fmt=CUR)
    set_cell(ws, "G17", "=G16-G15", fmt=CUR)

    set_cell(ws, "A18", "Ending Balance", bold=True)
    for col in YEARS_COLS:
        set_cell(ws, f"{col}18", f"={col}14-{col}17", fmt=CUR, bold=True)

    # --- Summary / covenant check ---
    section(ws, "A21", "Consolidated Debt Service & Coverage")
    set_cell(ws, "A22", "Total Debt Service (SBA P&I + Seller Note I+P)")
    for col in YEARS_COLS:
        set_cell(ws, f"{col}22", f"={col}7+{col}15+{col}17", fmt=CUR)

    set_cell(ws, "A23", "SDE / EBITDA (from Operating Model)")
    for col in YEARS_COLS:
        set_cell(ws, f"{col}23", f"='Operating Model'!{col}7", fmt=CUR)

    set_cell(ws, "A24", "DSCR (EBITDA / Total Debt Service)", bold=True)
    for col in YEARS_COLS:
        set_cell(ws, f"{col}24", f"={col}23/{col}22", fmt=MULT, bold=True,
                  note="Illustrative -- SBA/small-business lenders typically require "
                       "a minimum DSCR around 1.15x-1.25x. Confirm actual covenant "
                       "with the specific lender before relying on this.")

    set_cell(ws, "A25", "Free Cash Flow after Debt Service (to Equity)")
    for col in YEARS_COLS:
        set_cell(ws, f"{col}25", f"='Operating Model'!{col}12-{col}22", fmt=CUR)

    set_cell(ws, "A27", "Total Debt Outstanding, End of Year (SBA + Seller Note)", bold=True)
    for col in YEARS_COLS:
        set_cell(ws, f"{col}27", f"={col}10+{col}18", fmt=CUR, bold=True)

    return ws


def build_returns_analysis(wb):
    ws = wb.create_sheet("Returns Analysis")
    title(ws, "Returns Analysis")
    widen(ws, {"A": 44, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16, "G": 16})

    section(ws, "A3", "Base Case Exit")
    set_cell(ws, "A4", "Exit Year")
    set_cell(ws, "B4", "=Assumptions!B32", fmt="0")
    set_cell(ws, "A5", "Exit Multiple (x SDE)")
    set_cell(ws, "B5", "=Assumptions!B33", fmt=MULT)
    set_cell(ws, "A6", "Exit Year SDE / EBITDA")
    set_cell(ws, "B6", "=INDEX('Operating Model'!$C$7:$G$7,MATCH($B$4,'Operating Model'!$C$2:$G$2,0))", fmt=CUR)
    set_cell(ws, "A7", "Exit Enterprise Value")
    set_cell(ws, "B7", "=B6*B5", fmt=CUR)
    set_cell(ws, "A8", "Less: Net Debt at Exit")
    set_cell(ws, "B8", "=INDEX('Debt Schedule'!$C$27:$G$27,MATCH($B$4,'Operating Model'!$C$2:$G$2,0))", fmt=CUR)
    set_cell(ws, "A9", "Exit Equity Value", bold=True)
    set_cell(ws, "B9", "=B7-B8", fmt=CUR, bold=True)
    set_cell(ws, "A10", "Initial Sponsor Equity")
    set_cell(ws, "B10", "='Sources & Uses'!B11", fmt=CUR)
    set_cell(ws, "A11", "MOIC", bold=True)
    set_cell(ws, "B11", "=B9/B10", fmt=MULT, bold=True)
    set_cell(ws, "A12", "IRR (single lump-sum exit, no interim distributions)", bold=True)
    set_cell(ws, "B12", "=(B11)^(1/B4)-1", fmt=PCT, bold=True,
              note="Exact for this cash flow profile (one outflow at close, one "
                   "inflow at exit, nothing in between): IRR = MOIC^(1/years) - 1.")

    # --- Sensitivity: MOIC by exit multiple x exit year ---
    section(ws, "A15", "Sensitivity: MOIC by Exit Multiple x Exit Year")
    exit_years = [3, 4, 5]
    exit_multiples = [2.5, 3.0, 3.5]
    year_cols = {3: "C", 4: "D", 5: "E"}
    for j, yr in enumerate(exit_years):
        col = year_cols[yr]
        set_cell(ws, f"{col}16", yr, bold=True, fmt="0", fill=HEADER_FILL, font=HEADER_FONT)
    set_cell(ws, "B16", "Exit Multiple \\ Exit Year", bold=True)
    for i, mult in enumerate(exit_multiples):
        row = 17 + i
        set_cell(ws, f"B{row}", mult, font=BLUE, fmt=MULT, bold=True)
        for yr in exit_years:
            col = year_cols[yr]
            formula = (
                f"=(INDEX('Operating Model'!$C$7:$G$7,MATCH({col}$16,'Operating Model'!$C$2:$G$2,0))*$B{row}"
                f"-INDEX('Debt Schedule'!$C$27:$G$27,MATCH({col}$16,'Operating Model'!$C$2:$G$2,0)))"
                f"/'Sources & Uses'!$B$11"
            )
            set_cell(ws, f"{col}{row}", formula, fmt=MULT)

    # --- Sensitivity: IRR by exit multiple x exit year ---
    section(ws, "A22", "Sensitivity: IRR by Exit Multiple x Exit Year")
    for j, yr in enumerate(exit_years):
        col = year_cols[yr]
        set_cell(ws, f"{col}23", yr, bold=True, fmt="0", fill=HEADER_FILL, font=HEADER_FONT)
    set_cell(ws, "B23", "Exit Multiple \\ Exit Year", bold=True)
    for i, mult in enumerate(exit_multiples):
        row = 24 + i
        moic_row = 17 + i
        set_cell(ws, f"B{row}", mult, font=BLUE, fmt=MULT, bold=True)
        for yr in exit_years:
            col = year_cols[yr]
            set_cell(ws, f"{col}{row}", f"=({col}{moic_row})^(1/{col}$23)-1", fmt=PCT)

    return ws


def build_notes(wb):
    ws = wb.create_sheet("Notes & Limitations")
    title(ws, "Notes & Limitations")
    widen(ws, {"A": 100})
    notes = [
        "This model works at the pre-tax, pass-through SDE level appropriate for a "
        "small-business acquisition -- it does not model corporate income tax, "
        "depreciation/amortization schedules, or deferred tax effects.",
        "",
        "No cash sweep or voluntary prepayment is modeled on the SBA loan. Real SBA "
        "7(a) loans commonly carry prepayment penalties in years 1-3 (typically "
        "5%/3%/1% of the prepaid amount), so a cash sweep would need to weigh that "
        "cost -- omitted here for simplicity, not because it wouldn't matter.",
        "",
        "The exit calculation assumes a single lump-sum equity distribution at exit "
        "with no interim dividends or recapitalizations -- realistic for a "
        "buy-and-hold independent-sponsor deal, less realistic for a fund that "
        "takes dividend recaps along the way.",
        "",
        "SBA loan rate, sponsor equity %, and transaction cost assumptions are "
        "illustrative Q3 2026 ballparks, not sourced from a live SBA SOP or lender "
        "term sheet -- flagged individually on the Assumptions tab, consistent with "
        "this repo's verified/unverified labeling elsewhere.",
        "",
        "DSCR covenant threshold (1.15x-1.25x cited on the Debt Schedule tab) is a "
        "general small-business lending ballpark, not a specific lender's actual "
        "covenant -- confirm with the real lender's term sheet before relying on it.",
        "",
        "This is a portfolio project. The deal, business, and dollar figures are "
        "illustrative and tie to the same anonymized scenario as "
        "mna_valuation_engine.py elsewhere in this repo -- no real business's "
        "financial data appears here.",
        "",
        "Base-case MOIC/IRR (see Returns Analysis) run well above a typical "
        "institutional LBO. That is a direct, mechanical consequence of the SBA "
        "capital structure, not an optimistic growth assumption doing the work: "
        "sponsor equity is only ~10% of total uses (the SBA's real minimum "
        "injection for a change-of-ownership loan), so the same dollar of EBITDA "
        "growth and debt paydown is divided by a much smaller equity base than a "
        "traditional 30-40%-equity deal would use. This is the standard economic "
        "argument for SBA-financed search-fund and independent-sponsor "
        "acquisitions -- small equity checks amplify returns -- and it cuts both "
        "ways: the same leverage that produces an 8x+ MOIC here would produce a "
        "proportionally larger equity wipeout if SDE declined instead of grew.",
    ]
    for i, line in enumerate(notes, start=3):
        cell = ws[f"A{i}"]
        cell.value = line
        cell.font = Font(name=FONT_NAME, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def build():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_assumptions(wb)
    build_sources_uses(wb)
    build_operating_model(wb)
    build_debt_schedule(wb)
    build_returns_analysis(wb)
    build_notes(wb)
    wb.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    path = build()
    print(f"[SUCCESS] LBO model written -> {path}")
